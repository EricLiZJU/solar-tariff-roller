from solar_tariff_roller.models.project import ProjectProfile
from solar_tariff_roller.schemas.input import CalculationInput
from solar_tariff_roller.services.calc import (
    analyze_sensitivity,
    build_cashflow_result,
    estimate_initial_generation_10k_kwh,
    generate_sensitivity_values,
    project_generation_for_year,
    run_calculation,
)
from solar_tariff_roller.services.exporter import export_calculation_bundle
from solar_tariff_roller.services.parser import upsert_monthly_update
from solar_tariff_roller.services.parser.excel_reader import load_project_workbook
from solar_tariff_roller.services.solver import solve_tariff_by_target_irr
from solar_tariff_roller.services.solver import solve_latest_start_month_for_tariff

from openpyxl import Workbook
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from pathlib import Path
import json


def test_project_profile_can_be_created() -> None:
    project = ProjectProfile(project_name="demo", capacity_mwp=1.0)

    assert project.project_name == "demo"


def test_calculation_input_can_be_built() -> None:
    payload = CalculationInput(
        project={
            "project_name": "高宇液压分布式光伏",
            "capacity_mwp": 0.726635,
        },
        generation={
            "annual_sun_hours": 1329,
            "performance_ratio": 0.82,
        },
        consumption={
            "self_consumption_ratio": 0.8,
            "monthly_self_consumption_ratios": [0.8] * 12,
        },
        tariff={
            "feed_in_tariff": 0.4153,
            "consumer_tariff": 0.72,
            "consumer_discount_rate": 0.88,
        },
        cost={
            "capex_per_watt": 4.74,
            "total_investment_10k_cny": 344.83,
            "annual_om_10k_cny": 3.63,
        },
    )

    assert payload.project.project_name == "高宇液压分布式光伏"
    assert round(payload.discounted_consumer_tariff, 4) == 0.6336


def test_load_project_workbook_merges_two_excel_sources(tmp_path) -> None:
    calculation_path = tmp_path / "【测算表】测试项目.xlsx"
    station_path = tmp_path / "电站统计.xlsx"

    _build_calculation_workbook(calculation_path)
    _build_station_workbook(station_path)

    payload = load_project_workbook(calculation_path, station_path)

    assert payload.project.project_name == "测试项目"
    assert payload.project.station_name == "测试电站"
    assert payload.project.city == "台州"
    assert payload.project.capacity_mwp == 0.726635
    assert payload.generation.annual_sun_hours == 1329
    assert payload.tariff.feed_in_tariff == 0.4153
    assert payload.finance.discount_rate == 0.06
    assert payload.cost.total_investment_10k_cny == 344.42499
    assert len(payload.rolling.baseline_monthly_revenues_10k_cny) == 44
    assert len(payload.rolling.baseline_monthly_cashflows_10k_cny) == 301
    assert payload.rolling.annual_generation_forecast_10k_kwh[:3] == [77.21, 76.73, 76.26]
    assert payload.rolling.baseline_self_use_revenues_10k_cny[:2] == [39.13, 38.89]
    assert payload.rolling.baseline_feed_in_revenues_10k_cny[:2] == [6.41, 6.37]
    assert payload.rolling.baseline_discounted_consumer_tariff == 0.6336
    assert len(payload.monthly_records) == 12
    assert payload.monthly_records[0].period_label == "2024-01"
    assert payload.monthly_records[0].generation_10k_kwh == 10.0
    assert payload.monthly_records[-1].self_consumption_ratio == 0.91
    assert [round(value, 2) for value in payload.consumption.monthly_self_consumption_ratios] == [
        0.8,
        0.81,
        0.82,
        0.83,
        0.84,
        0.85,
        0.86,
        0.87,
        0.88,
        0.89,
        0.9,
        0.91,
    ]


def test_load_project_workbook_without_financial_sheet_still_works(tmp_path) -> None:
    calculation_path = tmp_path / "【测算表】测试项目.xlsx"
    station_path = tmp_path / "电站统计.xlsx"

    _build_calculation_workbook(calculation_path, include_financial_sheet=False)
    _build_station_workbook(station_path)

    payload = load_project_workbook(calculation_path, station_path)

    assert payload.project.project_name == "测试项目"
    assert payload.tariff.feed_in_tariff == 0.4153
    assert payload.finance.discount_rate == 0.06
    assert payload.rolling.irr_annualization_mode == "simple"
    assert len(payload.rolling.baseline_monthly_revenues_10k_cny) == 44


def test_load_project_workbook_handles_annual_sheet_variants(tmp_path) -> None:
    calculation_path = tmp_path / "【测算表】测试项目.xlsx"
    station_path = tmp_path / "电站统计.xlsx"

    _build_calculation_workbook(
        calculation_path,
        financial_sheet_name="现金流量（16%、10%）",
        financial_discount_rate=25.3769883,
    )
    _build_station_workbook(station_path)

    payload = load_project_workbook(calculation_path, station_path)

    assert payload.finance.discount_rate == 0.06
    assert payload.tax.annual_output_vat_rate == 0.16
    assert payload.tax.annual_capex_input_vat_primary_rate == 0.16
    assert payload.tax.annual_capex_input_vat_secondary_rate == 0.10


def test_load_project_workbook_parses_forecast_q_row_numbers(tmp_path) -> None:
    calculation_path = tmp_path / "【测算表】测试项目.xlsx"
    station_path = tmp_path / "电站统计.xlsx"

    _build_calculation_workbook(
        calculation_path,
        rolling_forecast_formulas={
            51: "=项目基础数据!Q66/12",
            52: "=项目基础数据!Q66/12",
            53: "=项目基础数据!Q67/12",
            54: "=项目基础数据!Q68/12",
        },
    )
    _build_station_workbook(station_path)

    payload = load_project_workbook(calculation_path, station_path)

    assert payload.rolling.forecast_q_row_numbers[:4] == [66, 66, 67, 68]


def test_load_project_workbook_applies_monthly_updates_and_refreshes_ratio(tmp_path) -> None:
    calculation_path = tmp_path / "【测算表】测试项目.xlsx"
    station_path = tmp_path / "电站统计.xlsx"
    monthly_update_dir = tmp_path / "monthly_updates"

    _build_calculation_workbook(calculation_path)
    _build_station_workbook(station_path)
    upsert_monthly_update(
        calculation_path,
        station_path,
        {
            "period_label": "2024-12",
            "generation_10k_kwh": 10.0,
            "self_consumed_10k_kwh": 9.5,
            "exported_10k_kwh": 0.5,
        },
        base_dir=monthly_update_dir,
    )

    payload = load_project_workbook(
        calculation_path,
        station_path,
        monthly_update_dir=monthly_update_dir,
    )

    assert payload.monthly_records[-1].period_label == "2024-12"
    assert payload.monthly_records[-1].self_consumed_10k_kwh == 9.5
    assert payload.monthly_records[-1].self_consumption_ratio == 0.95
    assert payload.consumption.monthly_self_consumption_ratios[-1] == 0.95
    assert round(payload.consumption.self_consumption_ratio, 6) == 0.714286


def test_first_pass_calculation_engine_builds_cashflow() -> None:
    payload = CalculationInput(
        project={
            "project_name": "高宇液压分布式光伏",
            "capacity_mwp": 0.726635,
            "operation_years": 3,
        },
        generation={
            "annual_sun_hours": 1329,
            "performance_ratio": 0.82,
            "first_year_degradation_pct": 2.5,
            "annual_degradation_pct": 0.6,
        },
        consumption={
            "self_consumption_ratio": 0.8,
        },
        tariff={
            "feed_in_tariff": 0.4153,
            "consumer_tariff": 0.72,
            "consumer_discount_rate": 0.88,
        },
        cost={
            "capex_per_watt": 4.74,
            "total_investment_10k_cny": 344.42499,
            "annual_rent_10k_cny": 0,
            "annual_om_10k_cny": 3.633175,
            "annual_insurance_10k_cny": 0.34442499,
        },
        finance={
            "discount_rate": 0.06,
        },
        rolling={
            "baseline_monthly_revenues_10k_cny": [3.8] * 24,
            "baseline_monthly_cashflows_10k_cny": [-344.42499] + [3.2] * 36,
            "baseline_self_use_revenues_10k_cny": [39.13, 38.89, 38.65],
            "baseline_feed_in_revenues_10k_cny": [6.41, 6.37, 6.33],
            "baseline_discounted_consumer_tariff": 0.6336,
            "historical_months_count": 24,
        },
    )

    assert round(estimate_initial_generation_10k_kwh(payload), 4) == 79.1872
    assert project_generation_for_year(payload, 1) == (77.21, 2.5)
    assert project_generation_for_year(payload, 2) == (76.73, 3.1)

    result = build_cashflow_result(payload)

    assert result.initial_generation_10k_kwh == 79.1872
    assert round(result.discounted_consumer_tariff, 4) == 0.6336
    assert result.monthly_irr is not None
    assert result.historical_months_count == 24
    assert len(result.monthly_projections) == 36
    assert len(result.annual_projections) == 3
    assert result.annual_projections[0].self_consumed_10k_kwh == 61.7688
    assert result.annual_projections[0].exported_10k_kwh == 15.4416
    assert result.annual_projections[0].gross_revenue_10k_cny > 40
    assert result.annual_projections[0].vat_credit_carry_10k_cny < 0
    assert result.project_npv_10k_cny < 0

    serialized = run_calculation(payload)
    assert serialized["historical_months_count"] == 24
    assert len(serialized["monthly_projections"]) == 36
    assert serialized["annual_projections"][0]["year"] == 1
    assert serialized["annual_projections"][1]["generation_10k_kwh"] == 76.73


def test_actual_months_replace_predicted_months_in_rolling_chain() -> None:
    payload = CalculationInput(
        project={
            "project_name": "滚动替换测试",
            "capacity_mwp": 0.726635,
            "operation_years": 3,
            "grid_connection_date": "2024-01-01",
        },
        generation={
            "annual_sun_hours": 1329,
            "performance_ratio": 0.82,
            "first_year_degradation_pct": 2.5,
            "annual_degradation_pct": 0.6,
        },
        consumption={
            "self_consumption_ratio": 0.8,
        },
        tariff={
            "feed_in_tariff": 0.4153,
            "consumer_tariff": 0.72,
            "consumer_discount_rate": 0.88,
        },
        cost={
            "capex_per_watt": 4.74,
            "total_investment_10k_cny": 344.42499,
            "annual_rent_10k_cny": 0,
            "annual_om_10k_cny": 3.633175,
        },
        monthly_records=[
            {
                "period_label": "2024-01",
                "generation_10k_kwh": 10.0,
                "self_consumed_10k_kwh": 8.0,
                "exported_10k_kwh": 2.0,
                "self_consumption_tariff": 0.701,
                "feed_in_tariff": 0.3678,
            }
        ],
        rolling={
            "baseline_monthly_revenues_10k_cny": [3.8] * 44,
            "baseline_monthly_cashflows_10k_cny": [-344.42499] + [3.2] * 36,
            "baseline_self_use_revenues_10k_cny": [39.13, 38.89, 38.65],
            "baseline_feed_in_revenues_10k_cny": [6.41, 6.37, 6.33],
            "baseline_discounted_consumer_tariff": 0.6336,
            "historical_months_count": 24,
        },
    )

    result = build_cashflow_result(payload)
    first_month = result.monthly_projections[0]

    assert first_month.period_type == "actual"
    assert first_month.generation_10k_kwh == 10.0
    assert first_month.self_consumed_10k_kwh == 8.0
    assert first_month.exported_10k_kwh == 2.0
    assert first_month.gross_revenue_10k_cny == round(8.0 * 0.701 + 2.0 * 0.3678, 8)


def test_monthly_update_can_fill_the_third_energy_value() -> None:
    from solar_tariff_roller.api.app import _build_monthly_record_input

    record = _build_monthly_record_input(
        period_label="2025-05",
        generation_10k_kwh=12.5,
        self_consumed_10k_kwh=9.2,
        exported_10k_kwh=None,
        self_consumption_tariff=0.688,
        feed_in_tariff=0.4021,
    )

    assert record.exported_10k_kwh == 3.3
    assert record.self_consumption_ratio == 0.736
    assert record.self_consumption_tariff == 0.688
    assert record.feed_in_tariff == 0.4021


def test_export_calculation_bundle_writes_json_and_excel(tmp_path) -> None:
    payload = CalculationInput(
        project={
            "project_name": "导出测试项目",
            "capacity_mwp": 0.726635,
            "operation_years": 25,
        },
        generation={
            "annual_sun_hours": 1329,
            "performance_ratio": 0.82,
        },
        consumption={
            "self_consumption_ratio": 0.8,
        },
        tariff={
            "feed_in_tariff": 0.4153,
            "consumer_tariff": 0.72,
            "consumer_discount_rate": 0.88,
        },
        cost={
            "capex_per_watt": 4.74,
            "total_investment_10k_cny": 344.42499,
            "annual_om_10k_cny": 3.633175,
            "annual_insurance_10k_cny": 0.34442499,
            "replacement_costs_10k_cny_by_year": {15: 13.392},
        },
        monthly_records=[
            {
                "period_label": "2024-01",
                "generation_10k_kwh": 10.0,
                "self_consumed_10k_kwh": 8.0,
                "exported_10k_kwh": 2.0,
                "self_consumption_ratio": 0.8,
            }
        ],
        rolling={
            "baseline_monthly_revenues_10k_cny": [3.8] * 44,
            "baseline_monthly_cashflows_10k_cny": [-344.42499] + [3.5] * 300,
            "baseline_self_use_revenues_10k_cny": [39.13 + i * 0.1 for i in range(25)],
            "baseline_feed_in_revenues_10k_cny": [6.41 + i * 0.05 for i in range(25)],
            "baseline_discounted_consumer_tariff": 0.6336,
        },
    )

    baseline_target = 0.09
    assert baseline_target is not None

    paths = export_calculation_bundle(
        payload,
        output_dir=tmp_path,
        stem="export_case",
        target_irr=baseline_target,
        sensitivity_parameter="tariff.consumer_tariff",
        sensitivity_start=0.68,
        sensitivity_stop=0.72,
        sensitivity_step=0.04,
    )

    assert paths["json"].exists()
    assert paths["excel"].exists()

    json_payload = json.loads(paths["json"].read_text(encoding="utf-8"))
    assert json_payload["project"]["project_name"] == "导出测试项目"
    assert len(json_payload["result"]["annual_projections"]) == 25
    assert len(json_payload["result"]["monthly_projections"]) == 300
    assert json_payload["target_irr_solution"]["target_irr"] == baseline_target
    assert json_payload["sensitivity_analysis"]["parameter_name"] == "tariff.consumer_tariff"
    assert len(json_payload["sensitivity_analysis"]["values"]) == 2
    assert len(json_payload["input"]["monthly_records"]) == 1

    workbook = load_workbook(paths["excel"], data_only=True)
    assert workbook.sheetnames == ["汇总", "年度汇总预览", "滚动月度现金流", "反算结果", "敏感性分析", "月度数据", "输入参数"]
    assert workbook["汇总"]["A1"].value == "滚动测算结果汇总"
    assert workbook["年度汇总预览"]["A2"].value == 1
    assert workbook["滚动月度现金流"]["A2"].value == 1
    assert workbook["敏感性分析"]["A1"].value == "敏感性分析结果"
    assert workbook["月度数据"]["A1"].value == "月度真实数据"


def test_solve_tariff_by_target_irr_recovers_current_tariff() -> None:
    payload = CalculationInput(
        project={
            "project_name": "高宇液压分布式光伏",
            "capacity_mwp": 0.726635,
            "operation_years": 25,
        },
        generation={
            "annual_sun_hours": 1329,
            "performance_ratio": 0.82,
            "first_year_degradation_pct": 2.5,
            "annual_degradation_pct": 0.6,
        },
        consumption={
            "self_consumption_ratio": 0.8,
        },
        tariff={
            "feed_in_tariff": 0.4153,
            "consumer_tariff": 0.72,
            "consumer_discount_rate": 0.88,
        },
        cost={
            "capex_per_watt": 4.74,
            "total_investment_10k_cny": 344.42499,
            "annual_rent_10k_cny": 0,
            "annual_om_10k_cny": 3.633175,
            "annual_insurance_10k_cny": 0.34442499,
            "replacement_costs_10k_cny_by_year": {15: 13.392},
        },
        finance={
            "discount_rate": 0.06,
        },
        rolling={
            "baseline_monthly_revenues_10k_cny": [3.8] * 44,
            "baseline_monthly_cashflows_10k_cny": [-344.42499] + [3.5] * 300,
            "baseline_self_use_revenues_10k_cny": [39.13 + i * 0.1 for i in range(25)],
            "baseline_feed_in_revenues_10k_cny": [6.41 + i * 0.05 for i in range(25)],
            "baseline_discounted_consumer_tariff": 0.6336,
        },
    )
    baseline_result = build_cashflow_result(payload)
    solved = solve_tariff_by_target_irr(payload, target_irr=0.09)

    assert baseline_result.project_irr is not None
    assert solved.solved_discounted_consumer_tariff > 0
    assert solved.solved_consumer_tariff > solved.solved_discounted_consumer_tariff
    assert abs(solved.solved_project_irr - 0.09) < 0.0001


def test_solve_latest_start_month_for_tariff_finds_latest_month_when_tariff_matches_base() -> None:
    payload = CalculationInput(
        project={
            "project_name": "启用月份测试",
            "capacity_mwp": 0.726635,
            "operation_years": 3,
        },
        generation={
            "annual_sun_hours": 1329,
            "performance_ratio": 0.82,
        },
        consumption={
            "self_consumption_ratio": 0.8,
        },
        tariff={
            "feed_in_tariff": 0.4153,
            "consumer_tariff": 0.72,
            "consumer_discount_rate": 0.88,
        },
        cost={
            "capex_per_watt": 4.74,
            "total_investment_10k_cny": 60.0,
            "annual_om_10k_cny": 0.5,
        },
        rolling={
            "baseline_monthly_revenues_10k_cny": [3.0] * 24,
            "baseline_monthly_cashflows_10k_cny": [-60.0] + [2.8] * 36,
            "baseline_self_use_revenues_10k_cny": [45.0, 44.5, 44.0],
            "baseline_feed_in_revenues_10k_cny": [7.0, 6.9, 6.8],
            "baseline_discounted_consumer_tariff": 0.6336,
            "historical_months_count": 24,
        },
    )

    baseline_result = build_cashflow_result(payload)
    assert baseline_result.project_irr is not None

    solved = solve_latest_start_month_for_tariff(
        payload,
        discounted_consumer_tariff=payload.discounted_consumer_tariff,
        target_irr=baseline_result.project_irr - 0.000001,
    )

    assert solved.reachable is True
    assert solved.earliest_feasible_start_month_index is not None
    assert solved.latest_feasible_start_month_index == 36
    assert solved.solved_project_irr is not None
    assert solved.solved_project_irr >= baseline_result.project_irr - 0.000001


def test_solve_latest_start_month_for_tariff_reports_unreachable_when_tariff_too_low() -> None:
    payload = CalculationInput(
        project={
            "project_name": "启用月份测试",
            "capacity_mwp": 0.726635,
            "operation_years": 3,
        },
        generation={
            "annual_sun_hours": 1329,
            "performance_ratio": 0.82,
        },
        consumption={
            "self_consumption_ratio": 0.8,
        },
        tariff={
            "feed_in_tariff": 0.4153,
            "consumer_tariff": 0.72,
            "consumer_discount_rate": 0.88,
        },
        cost={
            "capex_per_watt": 4.74,
            "total_investment_10k_cny": 60.0,
            "annual_om_10k_cny": 0.5,
        },
        rolling={
            "baseline_monthly_revenues_10k_cny": [3.0] * 24,
            "baseline_monthly_cashflows_10k_cny": [-60.0] + [2.8] * 36,
            "baseline_self_use_revenues_10k_cny": [45.0, 44.5, 44.0],
            "baseline_feed_in_revenues_10k_cny": [7.0, 6.9, 6.8],
            "baseline_discounted_consumer_tariff": 0.6336,
            "historical_months_count": 24,
        },
    )

    solved = solve_latest_start_month_for_tariff(
        payload,
        discounted_consumer_tariff=0.0,
        target_irr=0.45,
    )

    assert solved.reachable is False
    assert solved.earliest_feasible_start_month_index is None
    assert solved.latest_feasible_start_month_index is None
    assert "过低" in solved.message


def test_sensitivity_analysis_tracks_parameter_changes() -> None:
    payload = CalculationInput(
        project={
            "project_name": "敏感性测试项目",
            "capacity_mwp": 0.726635,
            "operation_years": 25,
        },
        generation={
            "annual_sun_hours": 1329,
            "performance_ratio": 0.82,
            "first_year_degradation_pct": 2.5,
            "annual_degradation_pct": 0.6,
        },
        consumption={
            "self_consumption_ratio": 0.8,
        },
        tariff={
            "feed_in_tariff": 0.4153,
            "consumer_tariff": 0.72,
            "consumer_discount_rate": 0.88,
        },
        cost={
            "capex_per_watt": 4.74,
            "total_investment_10k_cny": 344.42499,
            "annual_rent_10k_cny": 0,
            "annual_om_10k_cny": 3.633175,
            "annual_insurance_10k_cny": 0.34442499,
            "replacement_costs_10k_cny_by_year": {15: 13.392},
        },
        finance={
            "discount_rate": 0.06,
        },
        rolling={
            "baseline_monthly_revenues_10k_cny": [3.8] * 44,
            "baseline_monthly_cashflows_10k_cny": [-344.42499] + [3.5] * 300,
            "baseline_self_use_revenues_10k_cny": [39.13 + i * 0.1 for i in range(25)],
            "baseline_feed_in_revenues_10k_cny": [6.41 + i * 0.05 for i in range(25)],
            "baseline_discounted_consumer_tariff": 0.6336,
        },
    )

    values = generate_sensitivity_values(0.68, 0.76, 0.04)
    analysis = analyze_sensitivity(payload, "tariff.consumer_tariff", values)

    assert analysis.parameter_name == "tariff.consumer_tariff"
    assert len(analysis.points) == 3
    assert [point.parameter_value for point in analysis.points] == [0.68, 0.72, 0.76]
    assert analysis.points[0].project_npv_10k_cny < analysis.points[1].project_npv_10k_cny
    assert analysis.points[1].project_npv_10k_cny < analysis.points[2].project_npv_10k_cny
    assert analysis.points[0].discounted_consumer_tariff < analysis.points[2].discounted_consumer_tariff
    assert analysis.points[0].project_irr < analysis.points[1].project_irr


def test_web_page_focuses_on_desktop_workflow() -> None:
    from solar_tariff_roller.api.app import create_app

    client = TestClient(create_app())
    html = client.get("/").text

    assert "详细使用流程" not in html
    assert "<h2>分年现金流预览</h2>" not in html
    assert "<h2>结果文件位置</h2>" not in html
    assert "JSON API" not in html
    assert "Default Port" not in html
    assert "浏览文件" in html
    assert 'type="month"' in html
    assert "本月消纳电价" in html
    assert "本月上网电价" in html
    assert "三项电量里填任意两项即可" in html
    assert "固定折后消纳电价测算" in html
    assert "结果与图表下载" not in html
    assert "使用说明" in html
    assert "工作台" in html
    assert "tab-link active" in html


def test_help_page_contains_guidance_sections() -> None:
    from solar_tariff_roller.api.app import create_app

    client = TestClient(create_app())
    html = client.get("/help").text

    assert "详细使用流程" in html
    assert "填写注意事项" in html
    assert "结果说明" in html
    assert "工作台" in html
    assert "tab-link active" in html


def test_solve_page_shows_intermediate_calculation_sections(tmp_path) -> None:
    from solar_tariff_roller.api.app import create_app

    calculation_path = tmp_path / "【测算表】测试项目.xlsx"
    station_path = tmp_path / "电站统计.xlsx"
    _build_calculation_workbook(calculation_path)
    _build_station_workbook(station_path)

    client = TestClient(create_app())
    response = client.get(
        "/solve",
        params={
            "target_irr": 0.095977,
            "calculation_workbook": str(calculation_path),
            "station_workbook": str(station_path),
            "sensitivity_parameter": "tariff.consumer_tariff",
            "sensitivity_start": 0.68,
            "sensitivity_stop": 0.72,
            "sensitivity_step": 0.04,
        },
    )

    assert response.status_code == 200
    assert "计算中间过程" in response.text
    assert "滚动测算年度汇总预览" in response.text
    assert "滚动测算月度汇总预览" in response.text
    assert "项目概览与当前文件" in response.text
    assert "展开查看完整 25 个滚动年度汇总" in response.text
    assert "展开查看完整 300 个月滚动汇总" in response.text
    assert "真实替换" in response.text
    assert "只看真实" in response.text
    assert "filterPeriodRows" in response.text
    assert "<details class=\"accordion\"" in response.text
    assert "下载 Excel" in response.text
    assert "下载 NPV 图" in response.text
    assert "结果与图表下载" in response.text
    assert "class=\"sticky-table\"" in response.text


def test_post_solve_accepts_uploaded_workbooks_and_download_route(tmp_path) -> None:
    from solar_tariff_roller.api.app import create_app

    calculation_path = tmp_path / "【测算表】测试项目.xlsx"
    station_path = tmp_path / "电站统计.xlsx"
    _build_calculation_workbook(calculation_path)
    _build_station_workbook(station_path)

    client = TestClient(create_app())
    with calculation_path.open("rb") as calc_file, station_path.open("rb") as station_file:
        response = client.post(
            "/solve",
            data={
                "target_irr": "0.095977",
                "sensitivity_parameter": "tariff.consumer_tariff",
                "sensitivity_start": "0.68",
                "sensitivity_stop": "0.72",
                "sensitivity_step": "0.04",
                "calculation_workbook": "",
                "station_workbook": "",
            },
            files={
                "calculation_workbook_file": ("calc.xlsx", calc_file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                "station_workbook_file": ("station.xlsx", station_file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            },
        )

    assert response.status_code == 200
    assert "下载 Excel" in response.text
    assert "/download?path=" in response.text

    marker = "/download?path="
    start = response.text.find(marker)
    assert start != -1
    end = response.text.find('"', start)
    download_url = response.text[start:end]
    download_response = client.get(download_url)
    assert download_response.status_code == 200
    assert download_response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def _build_calculation_workbook(
    path,
    include_financial_sheet: bool = True,
    financial_sheet_name: str = "分年现金流量表及财务指标",
    financial_discount_rate: float = 0.06,
    rolling_forecast_formulas: dict[int, str] | None = None,
) -> None:
    workbook = Workbook()
    base = workbook.active
    base.title = "项目基础数据"
    financial = workbook.create_sheet(financial_sheet_name) if include_financial_sheet else None
    rolling = workbook.create_sheet("月滚动现金流量表")

    base["C5"] = 0.726635
    base["D5"] = 1329
    base["E5"] = 0.82
    base["D10"] = 4.74
    base["E10"] = 344.42499
    base["F10"] = 0
    base["G10"] = 3.633175
    base["C15"] = 0.8
    base["D20"] = 0
    base["D21"] = 0
    base["D22"] = 0
    base["D23"] = 0.4153
    base["D24"] = 0.72
    base["D25"] = 0.88
    base["D26"] = 0.6336
    base["D31"] = 2.5
    base["D32"] = 0.6
    generation_values = [77.21, 76.73, 76.26, 75.78, 75.31]
    while len(generation_values) < 25:
        generation_values.append(round(generation_values[-1] * 0.994, 2))
    for index, value in enumerate(generation_values[:25], start=31):
        base.cell(index, 7, value)
    for row_idx in range(63, 88):
        base.cell(row_idx, 13, round(6.41 - (row_idx - 63) * 0.04, 2))   # M
        base.cell(row_idx, 16, round(39.13 - (row_idx - 63) * 0.24, 2))  # P

    if financial is not None:
        financial["O5"] = financial_discount_rate
        financial["G7"] = 0.34442499
    rolling["F6"] = 344.42499
    for row_idx in range(7, 51):
        rolling.cell(row_idx, 3, 3.79583333333333)
    rolling.cell(50, 3, 3.79583333333333)
    if rolling_forecast_formulas:
        for row_idx, formula in rolling_forecast_formulas.items():
            rolling.cell(row_idx, 3, formula)
    for row_idx in range(6, 307):
        rolling.cell(row_idx, 13, 3.2 if row_idx > 6 else -344.42499)

    workbook.save(path)


def _build_station_workbook(path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"

    sheet.cell(4, 1, "台州")
    sheet.cell(4, 2, "测试电站")
    sheet.cell(4, 3, "台州")
    sheet.cell(4, 4, "临海市")
    sheet.cell(4, 5, 45292)
    sheet.cell(4, 7, 0.726635)
    sheet.cell(4, 457, 0.86)
    sheet.cell(4, 728, 0.69)

    _fill_monthly_block(sheet, start_col=44, values=[10.0 + index for index in range(12)])
    _fill_monthly_block(sheet, start_col=134, values=[8.0 + 0.5 * index for index in range(12)])
    _fill_monthly_block(sheet, start_col=224, values=[1.0 + 0.25 * index for index in range(12)])
    _fill_monthly_block(sheet, start_col=404, values=[0.8 + 0.01 * index for index in range(12)])

    workbook.save(path)


def _fill_monthly_block(sheet, start_col: int, values: list[float]) -> None:
    sheet.cell(3, start_col, 45292)
    for month_index in range(1, 12):
        sheet.cell(3, start_col + month_index, f"{month_index + 1}月")

    for offset, value in enumerate(values):
        sheet.cell(4, start_col + offset, value)
