"""JSON and Excel exporters for rolling calculation results."""

from __future__ import annotations

import json
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font

from solar_tariff_roller.schemas.input import CalculationInput
from solar_tariff_roller.services.calc import (
    analyze_sensitivity,
    build_cashflow_reconciliation,
    generate_sensitivity_values,
    run_calculation,
)
from solar_tariff_roller.services.solver import solve_tariff_by_target_irr


DEFAULT_EXPORT_DIR = Path("data/exports")


def export_calculation_bundle(
    payload: CalculationInput,
    output_dir: str | Path = DEFAULT_EXPORT_DIR,
    stem: str | None = None,
    reference_workbook_path: str | Path | None = None,
    target_irr: float | None = None,
    sensitivity_parameter: str | None = None,
    sensitivity_start: float | None = None,
    sensitivity_stop: float | None = None,
    sensitivity_step: float | None = None,
) -> dict[str, Path]:
    """Export one calculation result to both JSON and Excel files."""

    target_dir = Path(output_dir)
    target_dir.mkdir(parents=True, exist_ok=True)

    filename_stem = _build_export_stem(payload, stem)
    result = run_calculation(payload)
    export_payload = {
        "exported_at": datetime.now().isoformat(timespec="seconds"),
        "project": payload.project.model_dump(),
        "input": payload.model_dump(),
        "result": result,
    }
    effective_target_irr = target_irr if target_irr is not None else payload.finance.target_irr
    if effective_target_irr is not None:
        export_payload["target_irr_solution"] = asdict(
            solve_tariff_by_target_irr(payload, target_irr=effective_target_irr)
        )
    if sensitivity_parameter is not None:
        if None in (sensitivity_start, sensitivity_stop, sensitivity_step):
            raise ValueError("sensitivity range must include start, stop, and step")
        values = generate_sensitivity_values(sensitivity_start, sensitivity_stop, sensitivity_step)
        sensitivity = analyze_sensitivity(payload, sensitivity_parameter, values)
        export_payload["sensitivity_analysis"] = {
            "parameter_name": sensitivity.parameter_name,
            "values": [asdict(point) for point in sensitivity.points],
        }
    if reference_workbook_path is not None:
        export_payload["reconciliation"] = build_cashflow_reconciliation(payload, reference_workbook_path)

    json_path = target_dir / f"{filename_stem}.json"
    xlsx_path = target_dir / f"{filename_stem}.xlsx"

    _write_json_export(json_path, export_payload)
    _write_excel_export(xlsx_path, export_payload)

    return {
        "json": json_path,
        "excel": xlsx_path,
    }


def _write_json_export(path: Path, payload: dict[str, Any]) -> None:
    """Write a JSON export to disk."""

    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_excel_export(path: Path, payload: dict[str, Any]) -> None:
    """Write a structured Excel export with rolling cashflow and annual preview sheets."""

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "汇总"
    annual_sheet = workbook.create_sheet("年度汇总预览")
    rolling_sheet = workbook.create_sheet("滚动月度现金流")
    solve_sheet = workbook.create_sheet("反算结果") if "target_irr_solution" in payload else None
    sensitivity_sheet = workbook.create_sheet("敏感性分析") if "sensitivity_analysis" in payload else None
    monthly_sheet = workbook.create_sheet("月度数据") if payload["input"].get("monthly_records") else None
    input_sheet = workbook.create_sheet("输入参数")
    reconciliation_sheet = workbook.create_sheet("Excel对账") if "reconciliation" in payload else None

    _fill_summary_sheet(summary_sheet, payload)
    _fill_annual_sheet(annual_sheet, payload["result"]["annual_projections"])
    _fill_rolling_sheet(rolling_sheet, payload["result"].get("monthly_projections", []))
    if solve_sheet is not None:
        _fill_target_irr_sheet(solve_sheet, payload["target_irr_solution"])
    if sensitivity_sheet is not None:
        _fill_sensitivity_sheet(sensitivity_sheet, payload["sensitivity_analysis"])
    if monthly_sheet is not None:
        _fill_monthly_sheet(monthly_sheet, payload["input"]["monthly_records"])
    _fill_input_sheet(input_sheet, payload["input"])
    if reconciliation_sheet is not None:
        _fill_reconciliation_sheet(reconciliation_sheet, payload["reconciliation"])

    workbook.save(path)


def _fill_summary_sheet(sheet: Any, payload: dict[str, Any]) -> None:
    """Render key project metrics on the summary sheet."""

    project = payload["project"]
    result = payload["result"]

    rows = [
        ("项目名称", project["project_name"]),
        ("电站名称", project.get("station_name")),
        ("区域", project.get("region")),
        ("城市", project.get("city")),
        ("并网时间", project.get("grid_connection_date")),
        ("导出时间", payload["exported_at"]),
        ("", ""),
        ("初始年发电量(万kWh)", result["initial_generation_10k_kwh"]),
        ("折后用户侧电价(元/kWh)", result["discounted_consumer_tariff"]),
        ("滚动月度 IRR", result.get("monthly_irr")),
        ("初始投资流出(万元)", result["initial_outflow_10k_cny"]),
        ("资本开支进项税(万元)", result["capex_input_vat_10k_cny"]),
        ("项目净现值 NPV(万元)", result["project_npv_10k_cny"]),
        ("项目滚动年化 IRR", result["project_irr"]),
        ("累计净现金流(万元)", result["cumulative_cashflow_10k_cny"]),
        ("历史固化月份数", result.get("historical_months_count", 0)),
    ]
    if "target_irr_solution" in payload:
        solution = payload["target_irr_solution"]
        rows.extend(
            [
                ("", ""),
                ("目标 IRR", solution["target_irr"]),
                ("反算用户侧综合电价(元/kWh)", solution["solved_consumer_tariff"]),
                ("反算折后消纳电价(元/kWh)", solution["solved_discounted_consumer_tariff"]),
                ("反算校验 NPV(万元)", solution["solved_npv_10k_cny"]),
                ("反算校验滚动 IRR", solution["solved_project_irr"]),
            ]
        )
    if "sensitivity_analysis" in payload:
        sensitivity = payload["sensitivity_analysis"]
        rows.extend(
            [
                ("", ""),
                ("敏感性分析变量", sensitivity["parameter_name"]),
                ("敏感性分析点数", len(sensitivity["values"])),
            ]
        )
    monthly_records = payload["input"].get("monthly_records", [])
    if monthly_records:
        rows.extend(
            [
                ("", ""),
                ("月度数据条数", len(monthly_records)),
                ("最近月份", monthly_records[-1]["period_label"]),
            ]
        )

    sheet["A1"] = "滚动测算结果汇总"
    sheet["A1"].font = Font(bold=True, size=14)

    row_idx = 3
    for label, value in rows:
        sheet.cell(row_idx, 1, label)
        sheet.cell(row_idx, 2, value)
        if label:
            sheet.cell(row_idx, 1).font = Font(bold=True)
        row_idx += 1

    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 22


def _fill_annual_sheet(sheet: Any, annual_rows: list[dict[str, Any]]) -> None:
    """Render the annual preview sheet aggregated from rolling months."""

    headers = [
        "年份",
        "累计衰减(%)",
        "年发电量(万kWh)",
        "自用电量(万kWh)",
        "上网电量(万kWh)",
        "自用收益(万元)",
        "上网收益(万元)",
        "补贴收益(万元)",
        "总收入(万元)",
        "年度成本(万元)",
        "进项税(万元)",
        "销项税(万元)",
        "应缴增值税(万元)",
        "留抵税额(万元)",
        "附加税(万元)",
        "净现金流(万元)",
        "折现现金流(万元)",
        "累计现金流(万元)",
    ]
    keys = [
        "year",
        "degradation_pct",
        "generation_10k_kwh",
        "self_consumed_10k_kwh",
        "exported_10k_kwh",
        "self_consumption_revenue_10k_cny",
        "feed_in_revenue_10k_cny",
        "subsidy_revenue_10k_cny",
        "gross_revenue_10k_cny",
        "annual_cost_10k_cny",
        "input_vat_10k_cny",
        "output_vat_10k_cny",
        "vat_payable_10k_cny",
        "vat_credit_carry_10k_cny",
        "surcharge_tax_10k_cny",
        "net_cashflow_10k_cny",
        "discounted_cashflow_10k_cny",
        "cumulative_cashflow_10k_cny",
    ]

    for column_idx, header in enumerate(headers, start=1):
        sheet.cell(1, column_idx, header).font = Font(bold=True)

    for row_idx, row in enumerate(annual_rows, start=2):
        for column_idx, key in enumerate(keys, start=1):
            sheet.cell(row_idx, column_idx, row[key])

    for column_letter, width in {
        "A": 8,
        "B": 12,
        "C": 15,
        "D": 15,
        "E": 15,
        "F": 14,
        "G": 14,
        "H": 14,
        "I": 14,
        "J": 14,
        "K": 12,
        "L": 12,
        "M": 14,
        "N": 14,
        "O": 12,
        "P": 14,
        "Q": 14,
        "R": 14,
    }.items():
        sheet.column_dimensions[column_letter].width = width


def _fill_target_irr_sheet(sheet: Any, solution: dict[str, Any]) -> None:
    """Render target IRR reverse-solve results."""

    sheet["A1"] = "目标 IRR 滚动反算结果"
    sheet["A1"].font = Font(bold=True, size=14)

    rows = [
        ("目标 IRR", solution["target_irr"]),
        ("反算用户侧综合电价(元/kWh)", solution["solved_consumer_tariff"]),
        ("反算折后消纳电价(元/kWh)", solution["solved_discounted_consumer_tariff"]),
        ("反算校验 NPV(万元)", solution["solved_npv_10k_cny"]),
        ("反算校验滚动 IRR", solution["solved_project_irr"]),
    ]
    for row_idx, (label, value) in enumerate(rows, start=3):
        sheet.cell(row_idx, 1, label).font = Font(bold=True)
        sheet.cell(row_idx, 2, value)

    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 20


def _fill_rolling_sheet(sheet: Any, monthly_rows: list[dict[str, Any]]) -> None:
    """Render the rolling monthly cashflow detail sheet."""

    headers = [
        "月份序号",
        "运营年份",
        "年内月份",
        "阶段",
        "月发电量(万kWh)",
        "月自用电量(万kWh)",
        "月上网电量(万kWh)",
        "月含税收入(万元)",
        "年化收入基准(万元)",
        "保险费(万元)",
        "运维及租金(万元)",
        "更换成本(万元)",
        "成本合计(万元)",
        "进项税(万元)",
        "销项税(万元)",
        "增值税余额(万元)",
        "应缴增值税(万元)",
        "留抵税额(万元)",
        "附加税(万元)",
        "净现金流(万元)",
        "月折现系数",
        "折现现金流(万元)",
        "累计现金流(万元)",
    ]
    keys = [
        "month_index",
        "operating_year",
        "month_in_year",
        "period_type",
        "generation_10k_kwh",
        "self_consumed_10k_kwh",
        "exported_10k_kwh",
        "gross_revenue_10k_cny",
        "annualized_revenue_basis_10k_cny",
        "insurance_cost_10k_cny",
        "om_cost_10k_cny",
        "replacement_cost_10k_cny",
        "total_cost_10k_cny",
        "input_vat_10k_cny",
        "output_vat_10k_cny",
        "vat_balance_10k_cny",
        "vat_payable_10k_cny",
        "vat_credit_carry_10k_cny",
        "surcharge_tax_10k_cny",
        "net_cashflow_10k_cny",
        "discount_factor",
        "discounted_cashflow_10k_cny",
        "cumulative_cashflow_10k_cny",
    ]

    for column_idx, header in enumerate(headers, start=1):
        sheet.cell(1, column_idx, header).font = Font(bold=True)

    for row_idx, row in enumerate(monthly_rows, start=2):
        for column_idx, key in enumerate(keys, start=1):
            sheet.cell(row_idx, column_idx, row.get(key))

    for column_letter, width in {
        "A": 10,
        "B": 10,
        "C": 10,
        "D": 10,
        "E": 14,
        "F": 15,
        "G": 15,
        "H": 14,
        "I": 14,
        "J": 12,
        "K": 14,
        "L": 12,
        "M": 12,
        "N": 12,
        "O": 12,
        "P": 14,
        "Q": 14,
        "R": 14,
        "S": 12,
        "T": 14,
        "U": 12,
        "V": 14,
        "W": 14,
    }.items():
        sheet.column_dimensions[column_letter].width = width


def _fill_sensitivity_sheet(sheet: Any, sensitivity: dict[str, Any]) -> None:
    """Render one-variable sensitivity analysis results."""

    sheet["A1"] = "敏感性分析结果"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A3"] = "分析变量"
    sheet["B3"] = sensitivity["parameter_name"]
    sheet["A3"].font = Font(bold=True)

    headers = [
        "参数值",
        "NPV(万元)",
        "IRR",
        "累计现金流(万元)",
        "折后用户侧电价(元/kWh)",
    ]
    keys = [
        "parameter_value",
        "project_npv_10k_cny",
        "project_irr",
        "cumulative_cashflow_10k_cny",
        "discounted_consumer_tariff",
    ]

    header_row = 5
    for column_idx, header in enumerate(headers, start=1):
        sheet.cell(header_row, column_idx, header).font = Font(bold=True)

    for row_idx, row in enumerate(sensitivity["values"], start=header_row + 1):
        for column_idx, key in enumerate(keys, start=1):
            sheet.cell(row_idx, column_idx, row[key])

    for column_letter, width in {
        "A": 16,
        "B": 16,
        "C": 14,
        "D": 18,
        "E": 22,
    }.items():
        sheet.column_dimensions[column_letter].width = width

    if sensitivity["values"]:
        data_start_row = header_row + 1
        data_end_row = header_row + len(sensitivity["values"])

        line_chart = LineChart()
        line_chart.title = "NPV 敏感性折线图"
        line_chart.y_axis.title = "NPV(万元)"
        line_chart.x_axis.title = "参数值"
        line_chart.height = 8
        line_chart.width = 14
        npv_data = Reference(sheet, min_col=2, min_row=header_row, max_row=data_end_row)
        categories = Reference(sheet, min_col=1, min_row=data_start_row, max_row=data_end_row)
        line_chart.add_data(npv_data, titles_from_data=True)
        line_chart.set_categories(categories)
        sheet.add_chart(line_chart, "G5")

        bar_chart = BarChart()
        bar_chart.type = "col"
        bar_chart.style = 10
        bar_chart.title = "IRR 敏感性柱状图"
        bar_chart.y_axis.title = "IRR"
        bar_chart.x_axis.title = "参数值"
        bar_chart.height = 8
        bar_chart.width = 14
        irr_data = Reference(sheet, min_col=3, min_row=header_row, max_row=data_end_row)
        bar_chart.add_data(irr_data, titles_from_data=True)
        bar_chart.set_categories(categories)
        sheet.add_chart(bar_chart, "G22")


def _fill_input_sheet(sheet: Any, payload: dict[str, Any]) -> None:
    """Render normalized input parameters for traceability."""

    sheet["A1"] = "标准化输入参数"
    sheet["A1"].font = Font(bold=True, size=14)

    row_idx = 3
    for section_name, section_value in payload.items():
        sheet.cell(row_idx, 1, section_name).font = Font(bold=True)
        row_idx += 1

        if isinstance(section_value, list):
            for item in section_value:
                if isinstance(item, dict):
                    for key, value in item.items():
                        sheet.cell(row_idx, 2, key)
                        sheet.cell(row_idx, 3, value)
                        row_idx += 1
                    row_idx += 1
                else:
                    sheet.cell(row_idx, 2, item)
                    row_idx += 1
        elif isinstance(section_value, dict):
            for key, value in section_value.items():
                sheet.cell(row_idx, 2, key)
                if isinstance(value, (list, dict)):
                    sheet.cell(row_idx, 3, json.dumps(value, ensure_ascii=False))
                else:
                    sheet.cell(row_idx, 3, value)
                row_idx += 1
        else:
            sheet.cell(row_idx, 2, section_value)
            row_idx += 1

        row_idx += 1

    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 28
    sheet.column_dimensions["C"].width = 42


def _fill_monthly_sheet(sheet: Any, monthly_records: list[dict[str, Any]]) -> None:
    """Render merged monthly operating records."""

    sheet["A1"] = "月度真实数据"
    sheet["A1"].font = Font(bold=True, size=14)

    headers = [
        "月份",
        "发电量(万kWh)",
        "自用电量(万kWh)",
        "上网电量(万kWh)",
        "消纳率",
    ]
    keys = [
        "period_label",
        "generation_10k_kwh",
        "self_consumed_10k_kwh",
        "exported_10k_kwh",
        "self_consumption_ratio",
    ]

    header_row = 3
    for column_idx, header in enumerate(headers, start=1):
        sheet.cell(header_row, column_idx, header).font = Font(bold=True)

    for row_idx, row in enumerate(monthly_records, start=header_row + 1):
        for column_idx, key in enumerate(keys, start=1):
            sheet.cell(row_idx, column_idx, row.get(key))

    for column_letter, width in {
        "A": 14,
        "B": 18,
        "C": 18,
        "D": 18,
        "E": 12,
    }.items():
        sheet.column_dimensions[column_letter].width = width


def _build_export_stem(payload: CalculationInput, stem: str | None) -> str:
    """Build a safe filename stem for exports."""

    if stem:
        return stem

    raw_name = payload.project.project_name.strip()
    sanitized = "".join(char if char.isalnum() or char in ("-", "_") else "_" for char in raw_name)
    sanitized = sanitized.strip("_")
    if not sanitized:
        sanitized = "calculation_result"

    return f"{sanitized}_result"


def _fill_reconciliation_sheet(sheet: Any, reconciliation_rows: list[dict[str, Any]]) -> None:
    """Render column-by-column Excel/Python reconciliation."""

    headers = [
        "年份",
        "字段",
        "Excel",
        "Python",
        "差值",
    ]
    for column_idx, header in enumerate(headers, start=1):
        sheet.cell(1, column_idx, header).font = Font(bold=True)

    row_idx = 2
    for row in reconciliation_rows:
        year = row["year"]
        for key, value in row.items():
            if key == "year":
                continue
            sheet.cell(row_idx, 1, year)
            sheet.cell(row_idx, 2, key)
            sheet.cell(row_idx, 3, value["excel"])
            sheet.cell(row_idx, 4, value["python"])
            sheet.cell(row_idx, 5, value["diff"])
            row_idx += 1

    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 36
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 16
    sheet.column_dimensions["E"].width = 16
