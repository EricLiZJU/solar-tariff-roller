"""Reconciliation helpers between Python results and the source Excel cashflow sheet."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from solar_tariff_roller.schemas.input import CalculationInput
from solar_tariff_roller.services.calc.cashflow import build_cashflow_result


REFERENCE_SHEET = "分年现金流量表及财务指标"


def build_cashflow_reconciliation(
    payload: CalculationInput,
    workbook_path: str | Path,
) -> list[dict]:
    """Build a column-by-column reconciliation against the annual Excel cashflow sheet."""

    reference_rows = parse_reference_cashflow_sheet(workbook_path)
    result = build_cashflow_result(payload)

    python_rows = [_build_python_row_zero(result)] + [
        _build_python_year_row(row) for row in result.annual_projections
    ]

    comparison = []
    for excel_row, python_row in zip(reference_rows, python_rows, strict=True):
        row_comparison = {"year": excel_row["year"]}
        for key, excel_value in excel_row.items():
            if key == "year":
                continue
            python_value = python_row[key]
            diff = None
            if isinstance(excel_value, (int, float)) and isinstance(python_value, (int, float)):
                diff = round(python_value - excel_value, 8)
            row_comparison[key] = {
                "excel": excel_value,
                "python": python_value,
                "diff": diff,
            }
        comparison.append(row_comparison)

    return comparison


def parse_reference_cashflow_sheet(workbook_path: str | Path) -> list[dict]:
    """Parse the reference annual cashflow table from the source workbook."""

    workbook = load_workbook(workbook_path, data_only=True)
    sheet = workbook[REFERENCE_SHEET]
    rows = []
    for row_idx in range(6, 32):
        rows.append(
            {
                "year": sheet.cell(row_idx, 2).value,
                "gross_revenue_10k_cny": _number(sheet.cell(row_idx, 3).value),
                "revenue_excluding_vat_10k_cny": _number(sheet.cell(row_idx, 4).value),
                "output_vat_10k_cny": _number(sheet.cell(row_idx, 5).value),
                "construction_cost_10k_cny": _number(sheet.cell(row_idx, 6).value),
                "annual_insurance_10k_cny": _number(sheet.cell(row_idx, 7).value),
                "annual_om_and_rent_10k_cny": _number(sheet.cell(row_idx, 8).value),
                "annual_cost_10k_cny": _number(sheet.cell(row_idx, 9).value),
                "input_vat_10k_cny": _number(sheet.cell(row_idx, 10).value),
                "vat_payable_10k_cny": _number(sheet.cell(row_idx, 11).value),
                "surcharge_tax_10k_cny": _number(sheet.cell(row_idx, 12).value),
                "net_cashflow_10k_cny": _number(sheet.cell(row_idx, 13).value),
                "cumulative_cashflow_10k_cny": _number(sheet.cell(row_idx, 14).value),
                "present_value_factor": _number(sheet.cell(row_idx, 15).value),
                "discounted_cashflow_10k_cny": _number(sheet.cell(row_idx, 16).value),
                "cumulative_discounted_cashflow_10k_cny": _number(sheet.cell(row_idx, 17).value),
            }
        )
    return rows


def _build_python_row_zero(result) -> dict:
    return {
        "gross_revenue_10k_cny": 0.0,
        "revenue_excluding_vat_10k_cny": 0.0,
        "output_vat_10k_cny": 0.0,
        "construction_cost_10k_cny": result.initial_outflow_10k_cny,
        "annual_insurance_10k_cny": 0.0,
        "annual_om_and_rent_10k_cny": 0.0,
        "annual_cost_10k_cny": result.initial_outflow_10k_cny,
        "input_vat_10k_cny": result.capex_input_vat_10k_cny,
        "vat_payable_10k_cny": -result.capex_input_vat_10k_cny,
        "surcharge_tax_10k_cny": 0.0,
        "net_cashflow_10k_cny": -result.initial_outflow_10k_cny,
        "cumulative_cashflow_10k_cny": -result.initial_outflow_10k_cny,
        "present_value_factor": 1.0,
        "discounted_cashflow_10k_cny": -result.initial_outflow_10k_cny,
        "cumulative_discounted_cashflow_10k_cny": -result.initial_outflow_10k_cny,
    }


def _build_python_year_row(row) -> dict:
    return {
        "gross_revenue_10k_cny": row.gross_revenue_10k_cny,
        "revenue_excluding_vat_10k_cny": row.revenue_excluding_vat_10k_cny,
        "output_vat_10k_cny": row.output_vat_10k_cny,
        "construction_cost_10k_cny": row.construction_cost_10k_cny,
        "annual_insurance_10k_cny": row.annual_insurance_10k_cny,
        "annual_om_and_rent_10k_cny": row.annual_om_and_rent_10k_cny,
        "annual_cost_10k_cny": row.annual_cost_10k_cny,
        "input_vat_10k_cny": row.input_vat_10k_cny,
        "vat_payable_10k_cny": row.vat_payable_10k_cny if row.vat_payable_10k_cny > 0 else row.vat_credit_carry_10k_cny,
        "surcharge_tax_10k_cny": row.surcharge_tax_10k_cny,
        "net_cashflow_10k_cny": row.net_cashflow_10k_cny,
        "cumulative_cashflow_10k_cny": row.cumulative_cashflow_10k_cny,
        "present_value_factor": row.present_value_factor,
        "discounted_cashflow_10k_cny": row.discounted_cashflow_10k_cny,
        "cumulative_discounted_cashflow_10k_cny": row.cumulative_discounted_cashflow_10k_cny,
    }


def _number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return value
