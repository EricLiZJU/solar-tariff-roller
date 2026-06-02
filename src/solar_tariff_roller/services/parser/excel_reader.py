"""Excel parsers for the current project workbooks."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from solar_tariff_roller.schemas.input import CalculationInput, MonthlyGenerationRecordInput
from solar_tariff_roller.services.parser.monthly_updates import (
    DEFAULT_MONTHLY_UPDATE_DIR,
    calculate_recent_monthly_ratios,
    calculate_recent_self_consumption_ratio,
    load_monthly_updates,
    merge_monthly_records,
)


PROJECT_BASE_SHEET = "项目基础数据"
FINANCIAL_SHEET = "分年现金流量表及财务指标"
ROLLING_SHEET = "月滚动现金流量表"
STATION_SHEET = "Sheet1"


@dataclass(slots=True)
class ParsedStationData:
    """Station workbook data normalized for merging."""

    station_name: str | None
    region: str | None
    city: str | None
    district: str | None
    grid_connection_date: str | None
    capacity_mwp: float | None
    average_self_consumption_ratio: float | None
    feasibility_self_consumption_ratio: float | None
    average_consumer_tariff: float | None
    monthly_records: list[MonthlyGenerationRecordInput]


def load_project_workbook(
    calculation_workbook_path: str | Path,
    station_workbook_path: str | Path,
    monthly_update_dir: str | Path = DEFAULT_MONTHLY_UPDATE_DIR,
) -> CalculationInput:
    """Load both current Excel workbooks and merge them into one calculation input."""

    calculation_path = Path(calculation_workbook_path)
    station_path = Path(station_workbook_path)

    calculation_data = parse_calculation_workbook(calculation_path)
    station_data = parse_station_workbook(station_path)
    monthly_updates = load_monthly_updates(calculation_path, station_path, monthly_update_dir)
    has_monthly_updates = bool(monthly_updates)
    station_data.monthly_records = merge_monthly_records(station_data.monthly_records, monthly_updates)

    calculation_data["project"].update(
        {
            "station_name": station_data.station_name,
            "region": station_data.region,
            "city": station_data.city,
            "district": station_data.district,
            "grid_connection_date": station_data.grid_connection_date,
        }
    )

    if calculation_data["project"]["capacity_mwp"] <= 0 and station_data.capacity_mwp:
        calculation_data["project"]["capacity_mwp"] = station_data.capacity_mwp

    recent_self_consumption_ratio = calculate_recent_self_consumption_ratio(station_data.monthly_records)
    if has_monthly_updates and recent_self_consumption_ratio is not None:
        calculation_data["consumption"]["self_consumption_ratio"] = recent_self_consumption_ratio
    elif calculation_data["consumption"]["self_consumption_ratio"] == 0 and station_data.average_self_consumption_ratio:
        calculation_data["consumption"]["self_consumption_ratio"] = station_data.average_self_consumption_ratio

    calculation_data["consumption"]["feasibility_self_consumption_ratio"] = (
        station_data.feasibility_self_consumption_ratio
    )
    calculation_data["consumption"]["monthly_self_consumption_ratios"] = calculate_recent_monthly_ratios(
        station_data.monthly_records
    )

    if calculation_data["tariff"]["consumer_tariff"] == 0 and station_data.average_consumer_tariff is not None:
        calculation_data["tariff"]["consumer_tariff"] = station_data.average_consumer_tariff

    calculation_data["monthly_records"] = [record.model_dump() for record in station_data.monthly_records]
    return CalculationInput(**calculation_data)


def parse_calculation_workbook(workbook_path: str | Path) -> dict[str, Any]:
    """Parse the feasibility workbook into the standard calculation schema."""

    workbook = load_workbook(workbook_path, data_only=True)
    formula_workbook = load_workbook(workbook_path, data_only=False)
    base_sheet = workbook[PROJECT_BASE_SHEET]
    financial_sheet = _find_financial_sheet(workbook)
    rolling_sheet = workbook[ROLLING_SHEET] if ROLLING_SHEET in workbook.sheetnames else None
    rolling_formula_sheet = formula_workbook[ROLLING_SHEET] if ROLLING_SHEET in formula_workbook.sheetnames else None

    capacity_mwp = _as_float(base_sheet["C5"].value) or _as_float(base_sheet["C10"].value) or 0.0
    self_consumption_ratio = _as_float(base_sheet["C15"].value) or 0.0
    total_investment = (
        _as_float(rolling_sheet["F6"].value) if rolling_sheet is not None else None
    ) or _required_float(base_sheet["E10"].value, "项目基础数据!E10")

    discount_rate = None
    if financial_sheet is not None:
        discount_rate = _extract_discount_rate(financial_sheet)
    annual_output_vat_rate, annual_capex_primary_rate, annual_capex_secondary_rate = (
        _resolve_annual_tax_profile(financial_sheet.title) if financial_sheet is not None else (0.13, 0.13, 0.09)
    )

    return {
        "project": {
            "project_name": _normalize_project_name(Path(workbook_path).stem),
            "capacity_mwp": capacity_mwp,
            "operation_years": 25,
        },
        "generation": {
            "annual_sun_hours": _required_float(base_sheet["D5"].value, "项目基础数据!D5"),
            "performance_ratio": _required_float(base_sheet["E5"].value, "项目基础数据!E5"),
            "first_year_degradation_pct": _as_float(base_sheet["D31"].value) or 2.5,
            "annual_degradation_pct": _as_float(base_sheet["D32"].value) or 0.6,
        },
        "consumption": {
            "self_consumption_ratio": self_consumption_ratio,
            "monthly_self_consumption_ratios": [],
            "feasibility_self_consumption_ratio": None,
        },
        "tariff": {
            "feed_in_tariff": _required_float(base_sheet["D23"].value, "项目基础数据!D23"),
            "consumer_tariff": _required_float(base_sheet["D24"].value, "项目基础数据!D24"),
            "consumer_discount_rate": _required_float(base_sheet["D25"].value, "项目基础数据!D25"),
            "national_subsidy": _as_float(base_sheet["D20"].value) or 0.0,
            "provincial_subsidy": _as_float(base_sheet["D21"].value) or 0.0,
            "local_subsidy": _as_float(base_sheet["D22"].value) or 0.0,
        },
        "cost": {
            "capex_per_watt": _required_float(base_sheet["D10"].value, "项目基础数据!D10"),
            "total_investment_10k_cny": total_investment,
            "annual_rent_10k_cny": _as_float(base_sheet["F10"].value) or 0.0,
            "annual_om_10k_cny": _as_float(base_sheet["G10"].value) or 0.0,
            "annual_insurance_10k_cny": (
                _as_float(financial_sheet["G7"].value) or 0.0 if financial_sheet is not None else 0.0
            ),
            "replacement_costs_10k_cny_by_year": (
                _parse_replacement_costs(financial_sheet) if financial_sheet is not None else {}
            ),
        },
        "tax": {
            "output_vat_rate": 0.13,
            "input_vat_rate": 0.06,
            "surcharge_rate": 0.12,
            "annual_output_vat_rate": annual_output_vat_rate,
            "annual_capex_input_vat_primary_rate": annual_capex_primary_rate,
            "annual_capex_input_vat_secondary_rate": annual_capex_secondary_rate,
        },
        "finance": {
            "discount_rate": discount_rate or 0.06,
            "target_irr": None,
        },
        "rolling": {
            "baseline_monthly_revenues_10k_cny": _parse_rolling_baseline_revenues(rolling_sheet),
            "baseline_monthly_cashflows_10k_cny": _parse_rolling_baseline_cashflows(rolling_sheet),
            "annual_generation_forecast_10k_kwh": _parse_annual_generation_forecast(
                base_sheet,
                operation_years=25,
            ),
            "baseline_self_use_revenues_10k_cny": _parse_yearly_series(base_sheet, 63, 87, 16),
            "baseline_feed_in_revenues_10k_cny": _parse_yearly_series(base_sheet, 63, 87, 13),
            "baseline_discounted_consumer_tariff": _required_float(base_sheet["D26"].value, "项目基础数据!D26"),
            "historical_months_count": 44,
            "forecast_q_row_numbers": _parse_forecast_q_row_numbers(
                rolling_formula_sheet,
                historical_months_count=44,
                total_months=25 * 12,
            ),
            "irr_annualization_mode": "simple",
        },
        "monthly_records": [],
    }


def parse_station_workbook(workbook_path: str | Path) -> ParsedStationData:
    """Parse the historical station workbook and return normalized station data."""

    workbook = load_workbook(workbook_path, data_only=True)
    sheet = workbook[STATION_SHEET]
    data_row = _find_station_row(sheet)

    generation_start = _find_station_block_start(sheet, "每月发电量（万度）", fallback=44)
    self_consumed_start = _find_station_block_start(sheet, "每月消纳电量（万度）", fallback=134)
    exported_start = _find_station_block_start(sheet, "每月上网电量（万度）", fallback=224)
    ratio_start = _find_station_block_start(sheet, "每月平均消纳", fallback=404, header_row=2)

    generation_map = _parse_monthly_series(sheet, data_row, generation_start, generation_start + 89)
    self_consumed_map = _parse_monthly_series(sheet, data_row, self_consumed_start, self_consumed_start + 89)
    exported_map = _parse_monthly_series(sheet, data_row, exported_start, exported_start + 89)
    ratio_map = _parse_monthly_series(sheet, data_row, ratio_start, ratio_start + 89)

    labels = sorted(
        {
            *generation_map.keys(),
            *self_consumed_map.keys(),
            *exported_map.keys(),
            *ratio_map.keys(),
        }
    )
    monthly_records = [
        MonthlyGenerationRecordInput(
            period_label=label,
            generation_10k_kwh=generation_map.get(label),
            self_consumed_10k_kwh=self_consumed_map.get(label),
            exported_10k_kwh=exported_map.get(label),
            self_consumption_ratio=ratio_map.get(label),
        )
        for label in labels
    ]

    return ParsedStationData(
        station_name=_as_str(sheet.cell(data_row, 2).value),
        region=_as_str(sheet.cell(data_row, 1).value),
        city=_as_str(sheet.cell(data_row, 3).value),
        district=_as_str(sheet.cell(data_row, 4).value),
        grid_connection_date=_excel_date_string(sheet.cell(data_row, 5).value),
        capacity_mwp=_as_float(sheet.cell(data_row, 7).value),
        average_self_consumption_ratio=_as_float(sheet.cell(data_row, 457).value),
        feasibility_self_consumption_ratio=_as_float(sheet.cell(data_row, 458).value),
        average_consumer_tariff=_as_float(sheet.cell(data_row, 728).value),
        monthly_records=monthly_records,
    )


def _find_station_block_start(sheet: Any, title: str, fallback: int, header_row: int = 3) -> int:
    """Find the first column for a station workbook metric block by header title."""

    for col_idx in range(1, sheet.max_column + 1):
        value = sheet.cell(header_row, col_idx).value
        if isinstance(value, str) and value.strip() == title:
            return col_idx
    return fallback


def _find_station_row(sheet: Any) -> int:
    """Find the first data row in the station statistics workbook."""

    for row_idx in range(4, sheet.max_row + 1):
        station_name = sheet.cell(row_idx, 2).value
        region = sheet.cell(row_idx, 1).value
        if isinstance(station_name, str) and station_name.strip() == "电站名称":
            continue
        if isinstance(region, str) and region.strip() == "序号":
            continue

        capacity = _as_float(sheet.cell(row_idx, 6).value) or _as_float(sheet.cell(row_idx, 7).value)
        has_monthly_value = any(_as_float(sheet.cell(row_idx, col_idx).value) is not None for col_idx in (7, 97, 187, 367))

        if (station_name not in (None, "") or region not in (None, "")) and (capacity is not None or has_monthly_value):
            return row_idx

    raise ValueError("No station data row found in station workbook")


def _parse_monthly_series(sheet: Any, row_idx: int, start_col: int, end_col: int) -> dict[str, float]:
    """Parse one monthly metric block into a mapping keyed by YYYY-MM."""

    current_year: int | None = None
    result: dict[str, float] = {}

    for col_idx in range(start_col, end_col + 1):
        row4_header = sheet.cell(4, col_idx).value
        row3_header = sheet.cell(3, col_idx).value
        header = row4_header if _looks_like_period_header(row4_header) else row3_header
        current_year, period_label = _resolve_period_label(header, current_year)
        if period_label is None:
            continue

        value = _as_float(sheet.cell(row_idx, col_idx).value)
        if value is None:
            continue

        result[period_label] = value

    return result


def _looks_like_period_header(value: Any) -> bool:
    """Check whether a cell looks like a month or year header."""

    if isinstance(value, datetime):
        return True

    if isinstance(value, (int, float)) and value > 30000:
        return True

    if isinstance(value, str):
        stripped = value.strip()
        if re.fullmatch(r"(20\d{2})年", stripped):
            return True
        if re.fullmatch(r"(\d{1,2})月", stripped):
            return True

    return False


def _resolve_period_label(header_value: Any, current_year: int | None) -> tuple[int | None, str | None]:
    """Convert the workbook header cell to a normalized YYYY-MM label."""

    if isinstance(header_value, datetime):
        return header_value.year, f"{header_value.year:04d}-{header_value.month:02d}"

    if isinstance(header_value, (int, float)) and header_value > 30000:
        converted = from_excel(header_value)
        return converted.year, f"{converted.year:04d}-{converted.month:02d}"

    if isinstance(header_value, str):
        year_match = re.fullmatch(r"(20\d{2})年", header_value.strip())
        if year_match:
            return int(year_match.group(1)), None

        month_match = re.fullmatch(r"(\d{1,2})月", header_value.strip())
        if month_match and current_year is not None:
            month = int(month_match.group(1))
            return current_year, f"{current_year:04d}-{month:02d}"

    return current_year, None


def _normalize_project_name(file_stem: str) -> str:
    """Convert the workbook file name into a cleaner project name."""

    return re.sub(r"^【[^】]+】", "", file_stem).strip()


def _parse_rolling_baseline_revenues(sheet: Any | None) -> list[float]:
    """Read the fixed historical monthly revenues from the rolling sheet."""

    if sheet is None:
        return []

    values: list[float] = []
    for row_idx in range(7, 51):
        value = _as_float(sheet.cell(row_idx, 3).value)
        values.append(value or 0.0)
    return values


def _parse_rolling_baseline_cashflows(sheet: Any | None) -> list[float]:
    """Read the monthly cashflow baseline used to validate the current IRR."""

    if sheet is None:
        return []

    values: list[float] = []
    for row_idx in range(6, 307):
        value = _as_float(sheet.cell(row_idx, 13).value)
        values.append(value or 0.0)
    return values


def _parse_forecast_q_row_numbers(
    sheet: Any | None,
    historical_months_count: int,
    total_months: int,
) -> list[int]:
    """Parse forecast-month Q-row references from rolling-sheet C-column formulas."""

    if sheet is None:
        return []

    start_row = 7 + historical_months_count
    end_row = 6 + total_months
    refs: list[int] = []
    for row_idx in range(start_row, end_row + 1):
        value = sheet.cell(row_idx, 3).value
        if not isinstance(value, str):
            continue
        match = re.search(r"Q(\d+)\s*/\s*12", value)
        if match is not None:
            refs.append(int(match.group(1)))
    return refs


def _parse_annual_generation_forecast(base_sheet: Any, operation_years: int) -> list[float]:
    """Read the annual generation forecast block when present."""

    values: list[float] = []
    for row_idx in range(31, 31 + operation_years):
        value = _as_float(base_sheet.cell(row_idx, 7).value)
        if value is None:
            break
        values.append(round(value, 2))
    return values


def _parse_yearly_series(sheet: Any, start_row: int, end_row: int, column_idx: int) -> list[float]:
    """Read one yearly value block from a fixed column."""

    values: list[float] = []
    for row_idx in range(start_row, end_row + 1):
        value = _as_float(sheet.cell(row_idx, column_idx).value)
        if value is not None:
            values.append(float(value))
    return values


def _excel_date_string(value: Any) -> str | None:
    """Convert Excel serial dates or datetime values to ISO strings."""

    if isinstance(value, datetime):
        return value.date().isoformat()

    if isinstance(value, (int, float)) and value > 30000:
        return from_excel(value).date().isoformat()

    return _as_str(value)


def _required_float(value: Any, label: str) -> float:
    """Read a required numeric value from a workbook cell."""

    number = _as_float(value)
    if number is None:
        raise ValueError(f"Expected numeric value at {label}")

    return number


def _as_float(value: Any) -> float | None:
    """Convert workbook values to floats, skipping blanks and placeholder marks."""

    if value in (None, "", "——", "-", "--"):
        return None

    if isinstance(value, bool):
        return float(value)

    if isinstance(value, (int, float)):
        return float(value)

    if isinstance(value, str):
        cleaned = value.strip().replace(",", "")
        if cleaned in ("", "——", "-", "--"):
            return None
        try:
            return float(cleaned)
        except ValueError:
            return None

    return None


def _as_str(value: Any) -> str | None:
    """Convert workbook cell contents to a stripped string."""

    if value in (None, ""):
        return None

    return str(value).strip()


def _find_financial_sheet(workbook: Any) -> Any | None:
    """Find the annual cashflow sheet across workbook naming variants."""

    if FINANCIAL_SHEET in workbook.sheetnames:
        return workbook[FINANCIAL_SHEET]

    for sheet_name in workbook.sheetnames:
        if sheet_name == ROLLING_SHEET:
            continue
        if "现金流量" in sheet_name:
            return workbook[sheet_name]

    return None


def _resolve_annual_tax_profile(sheet_name: str) -> tuple[float, float, float]:
    """Infer annual summary tax rates from the sheet naming convention."""

    if "16%、10%" in sheet_name:
        return 0.16, 0.16, 0.10
    if "13%、9%" in sheet_name:
        return 0.13, 0.13, 0.09
    return 0.13, 0.13, 0.09


def _extract_discount_rate(financial_sheet: Any) -> float | None:
    """Extract a plausible annual discount rate from known workbook layouts."""

    for cell_ref in ("O5", "N3"):
        value = _as_float(financial_sheet[cell_ref].value)
        if value is not None and 0 < value <= 1:
            return value
    return None


def _parse_replacement_costs(financial_sheet: Any) -> dict[int, float]:
    """Parse year-specific replacement construction costs from the financial sheet."""

    replacements: dict[int, float] = {}
    for row_idx in range(7, 32):
        year = financial_sheet.cell(row_idx, 2).value
        construction_cost = _as_float(financial_sheet.cell(row_idx, 6).value)
        if isinstance(year, (int, float)) and construction_cost and construction_cost > 0:
            replacements[int(year)] = construction_cost

    return replacements
